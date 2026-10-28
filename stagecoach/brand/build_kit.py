"""Build the Stagecoach brand kit: logos, colour, type, print and digital assets.

    python3 brand/build_kit.py OUTPUT_DIR

Needs: reportlab, fonttools, shapely, python-docx, Pillow, Node with Playwright
(for PNG rendering), and pdftocairo (for previews).
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import logo as LG  # noqa: E402

from reportlab.graphics.shapes import Drawing  # noqa: E402
from reportlab.graphics import renderPDF, renderPS  # noqa: E402

OUT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(HERE, 'kit')
INCH = 72.0
FONTS = os.path.join(HERE, 'fonts')
SITE = os.path.dirname(HERE)

LOGOS = [
    # folder, kind, tagline, nominal width or height (inches) for the PDF/EPS master, png widths
    ('Horizontal', 'horizontal', True, ('w', 6.0), [600, 1600, 4000]),
    ('Horizontal_NoTagline', 'horizontal', False, ('w', 6.0), [600, 1600, 4000]),
    ('Stacked', 'stacked', True, ('w', 4.0), [500, 1200, 3000]),
    ('Symbol', 'symbol', True, ('h', 2.0), [256, 1024, 3000]),
    ('Wordmark', 'wordmark', True, ('w', 6.0), [600, 1600, 4000]),
]


def p(*parts):
    path = os.path.join(OUT, *parts)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    return path


# ---------------------------------------------------------------- PNG rendering
class PNG:
    """Queue SVG to PNG jobs and render them in one Playwright run."""
    jobs = []

    @classmethod
    def add(cls, svg, path, width=None, height=None, bg=None):
        cls.jobs.append({'svg': svg, 'path': path, 'w': width, 'h': height, 'bg': bg})

    @classmethod
    def add_html(cls, html, path, width, height, scale=1):
        cls.jobs.append({'html': html, 'path': path, 'w': width, 'h': height, 'scale': scale})

    @classmethod
    def run(cls):
        if not cls.jobs:
            return
        tmp = tempfile.mkdtemp()
        jf = os.path.join(tmp, 'jobs.json')
        json.dump(cls.jobs, open(jf, 'w'))
        js = os.path.join(tmp, 'r.js')
        open(js, 'w').write(r'''
const { chromium } = require('playwright'); const fs = require('fs');
const jobs = JSON.parse(fs.readFileSync(process.argv[2], 'utf8'));
(async () => {
  const b = await chromium.launch();
  for (const j of jobs) {
    if (j.html) {
      const p = await b.newPage({ viewport: { width: j.w, height: j.h }, deviceScaleFactor: j.scale || 1 });
      await p.setContent(j.html, { waitUntil: 'networkidle' }); await p.evaluate(() => document.fonts.ready);
      await p.screenshot({ path: j.path, type: j.path.endsWith('.jpg') ? 'jpeg' : 'png', quality: j.path.endsWith('.jpg') ? 90 : undefined });
      await p.close(); continue;
    }
    const vb = j.svg.match(/viewBox="([^"]+)"/)[1].split(/\s+/).map(Number);
    const ar = vb[3] / vb[2];
    const w = j.w || Math.round(j.h / ar), h = j.h || Math.round(j.w * ar);
    const p = await b.newPage({ viewport: { width: w, height: h } });
    const svg = j.svg.replace(/width="[^"]*" height="[^"]*"/, `width="${w}" height="${h}"`);
    await p.setContent(`<html><body style="margin:0;background:${j.bg || 'transparent'}">${svg}</body></html>`);
    await p.screenshot({ path: j.path, omitBackground: !j.bg, clip: { x: 0, y: 0, width: w, height: h } });
    await p.close();
  }
  await b.close();
})();''')
        subprocess.run(['node', js, jf], check=True, cwd=tmp, env={**os.environ, 'NODE_PATH': subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True).stdout.strip()})
        cls.jobs = []


# ---------------------------------------------------------------- logo files
def inner_svg(L, cw):
    """The logo's <path> elements only, origin at the top-left of the ink, for embedding."""
    import re
    s = LG.to_svg(L, cw, pad=0)
    s = re.sub(r'^<svg[^>]*>\s*<title>.*?</title>\s*', '', s, flags=re.S)
    return s.rsplit('</svg>', 1)[0]


def logo_drawing(L, cw, size_pts, axis, pad_units=None, cmyk=True):
    pad = LG.clear_space(L) if pad_units is None else pad_units
    k = size_pts / (L['w'] if axis == 'w' else L['h'])
    W, H = (L['w'] + 2 * pad) * k, (L['h'] + 2 * pad) * k
    d = Drawing(W, H)
    for s in LG.rl_paths(L, cw, pad * k, H - pad * k, k, cmyk=cmyk):
        d.add(s)
    return d


def logo_pdf(L, cw, size_pts, axis, path):
    """Vector CMYK logo on a transparent page sized to the logo plus clear space. No fonts referenced."""
    pad = LG.clear_space(L)
    k = size_pts / (L['w'] if axis == 'w' else L['h'])
    W, H = (L['w'] + 2 * pad) * k, (L['h'] + 2 * pad) * k
    c = rl_canvas.Canvas(path, pagesize=(W, H), initialFontName='Montserrat-Regular', initialFontSize=10)
    c.setTitle(os.path.splitext(os.path.basename(path))[0].replace('_', ' '))
    c.setAuthor('Stagecoach Concierge Emergency Medical Services')
    draw_logo(c, L, cw, pad * k, H - pad * k, width=L['w'] * k)
    c.showPage()
    c.save()


def build_logos():
    for folder, kind, tag, (axis, inches), widths in LOGOS:
        L = LG.layout(kind, tag)
        for cw in LG.COLOURWAYS:
            base = f'Stagecoach_{folder}_{cw}'
            svg = LG.to_svg(L, cw)
            open(p('01_Logo', folder, 'SVG', base + '.svg'), 'w').write(svg)
            d = logo_drawing(L, cw, inches * INCH, axis)
            logo_pdf(L, cw, inches * INCH, axis, p('01_Logo', folder, 'PDF_print', base + '.pdf'))
            renderPS.drawToFile(d, p('01_Logo', folder, 'EPS_print', base + '.eps'))
            for wpx in widths:
                if kind == 'symbol':
                    PNG.add(svg, p('01_Logo', folder, 'PNG', f'{base}_{wpx}px.png'), height=wpx)
                else:
                    PNG.add(svg, p('01_Logo', folder, 'PNG', f'{base}_{wpx}px.png'), width=wpx)
    # app icon: cream S on navy, full bleed square (platforms round the corners)
    L = LG.layout('symbol')
    pad = (L['h'] * 1.62 - L['w']) / 2
    sq = L['h'] * 1.62
    icon_svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {sq:.2f} {sq:.2f}" width="{sq:.2f}" height="{sq:.2f}">'
                f'<rect width="100%" height="100%" fill="{LG.COLOURS["navy"]["hex"]}"/>'
                f'<g transform="translate({pad:.2f} {(sq - L["h"]) / 2:.2f})">{inner_svg(L, "Reverse")}</g></svg>')
    open(p('01_Logo', 'App_Icon', 'Stagecoach_AppIcon.svg'), 'w').write(icon_svg)
    for px in (180, 512, 1024):
        PNG.add(icon_svg, p('01_Logo', 'App_Icon', f'Stagecoach_AppIcon_{px}px.png'), width=px)


# ---------------------------------------------------------------- shared print helpers
from reportlab.pdfgen import canvas as rl_canvas  # noqa: E402
from reportlab.pdfbase import pdfmetrics  # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont as RLTTFont  # noqa: E402
from reportlab.lib.colors import CMYKColor  # noqa: E402

for _w in ('Regular', 'Medium', 'SemiBold', 'Bold'):
    pdfmetrics.registerFont(RLTTFont(f'Montserrat-{_w}', os.path.join(FONTS, f'Montserrat-{_w}.ttf')))


def cmyk(role):
    return CMYKColor(*[v / 100 for v in LG.COLOURS[role]['cmyk']])


def draw_logo(c, L, cw, x, y_top, width=None, height=None):
    """Draw the outlined logo with the top-left of its ink at (x, y_top), sized by width or height.
    Paths go straight onto the canvas, so no fonts are referenced."""
    k = (width / L['w']) if width else (height / L['h'])
    roles = LG.COLOURWAYS[cw]
    for key, role, mode in (('s', 's', 0), ('dot', 'dot', 1), ('text', 'text', 1)):
        ops = L[key]
        if not ops:
            continue
        c.setFillColor(cmyk(roles[role]))
        path = c.beginPath()
        for op in ops:
            if op[0] == 'M':
                path.moveTo(x + op[1] * k, y_top - op[2] * k)
            elif op[0] == 'L':
                path.lineTo(x + op[1] * k, y_top - op[2] * k)
            elif op[0] == 'C':
                v = op[1:]
                path.curveTo(*[x + v[i] * k if i % 2 == 0 else y_top - v[i] * k for i in range(6)])
            else:
                path.close()
        c.drawPath(path, stroke=0, fill=1, fillMode=mode)
    return L['w'] * k, L['h'] * k


def tracked(c, text, x, y, font, size, track_em=0.0, colour=None, align='left'):
    if colour is not None:
        c.setFillColor(colour)
    w = pdfmetrics.stringWidth(text, font, size) + track_em * size * (len(text) - 1)
    if align == 'center':
        x -= w / 2
    elif align == 'right':
        x -= w
    t = c.beginText(x, y)
    t.setFont(font, size)
    t.setCharSpace(track_em * size)
    t.textOut(text)
    c.drawText(t)
    return w


def print_page(path, trim_w, trim_h, bleed=0.125 * INCH):
    """A canvas whose page is trim plus bleed, with TrimBox and BleedBox set for the printer."""
    c = rl_canvas.Canvas(path, pagesize=(trim_w + 2 * bleed, trim_h + 2 * bleed), initialFontName='Montserrat-Regular', initialFontSize=10)
    c.setTitle(os.path.splitext(os.path.basename(path))[0].replace('_', ' '))
    c.setAuthor('Stagecoach Concierge Emergency Medical Services')
    c.setTrimBox((bleed, bleed, bleed + trim_w, bleed + trim_h))
    c.setBleedBox((0, 0, trim_w + 2 * bleed, trim_h + 2 * bleed))
    return c, bleed


# ---------------------------------------------------------------- business cards
TEAM = [
    ('Noah Gray, EMT', 'Co-Founder & CEO'),
    ('David Vitberg, MD', 'Co-Founder & Chief Medical Officer'),
    ('Waleed Haleem, PMP', 'Co-Founder & Chief Operating Officer'),
    ('Scott Sehon, MBA', 'Co-Founder & Chief Financial Officer'),
    ('Liza White', 'VP of Strategy & Partnerships'),
    ('Ameen Elhamdani, NRP', 'EMS Operations Manager'),
    ('Gupi D. Howie, PA-C', 'Director of Advanced Practice Providers'),
    ('Yaffa Vitberg, MD', 'Pediatric Medical Director'),
    ('Eve Hill, RN, NRP', 'Quality Improvement Officer'),
    ('Charles White, DDS', 'Emergency Dental Response Partner'),
]
PHONE, EMAIL, WEB = '202.839.9439', 'medical@thestagecoachgroup.com', 'stagecoachmedical.com'
ADDRESS = '3299 K St NW #210, Washington, DC 20007'
CARD_W, CARD_H = 3.5 * INCH, 2.0 * INCH


def card_front(c, b):
    c.setFillColor(cmyk('navy'))
    c.rect(0, 0, CARD_W + 2 * b, CARD_H + 2 * b, stroke=0, fill=1)
    L = LG.layout('horizontal', True)
    w = 2.45 * INCH
    h = L['h'] * w / L['w']
    draw_logo(c, L, 'Reverse', b + (CARD_W - w) / 2, b + (CARD_H + h) / 2, width=w)


def card_back(c, b, name, role):
    c.setFillColor(cmyk('cream'))
    c.rect(0, 0, CARD_W + 2 * b, CARD_H + 2 * b, stroke=0, fill=1)
    m = 0.25 * INCH
    x, top = b + m, b + CARD_H - m
    # name and role
    tracked(c, name, x, top - 10, 'Montserrat-Bold', 10.5, 0, cmyk('navy'))
    role_caps = role.upper()
    size, track = 6.4, 0.14
    maxw = CARD_W - 2 * m - 0.42 * INCH
    lines = [role_caps]
    if pdfmetrics.stringWidth(role_caps, 'Montserrat-SemiBold', size) + track * size * len(role_caps) > maxw:
        words, lines, cur = role_caps.split(), [], ''
        for wd in words:
            t = (cur + ' ' + wd).strip()
            if pdfmetrics.stringWidth(t, 'Montserrat-SemiBold', size) + track * size * len(t) > maxw and cur:
                lines.append(cur)
                cur = wd
            else:
                cur = t
        lines.append(cur)
    y = top - 10 - 12.5
    for ln in lines:
        tracked(c, ln, x, y, 'Montserrat-SemiBold', size, track, cmyk('coral_ink'))
        y -= 9
    # symbol, top right
    S = LG.layout('symbol')
    draw_logo(c, S, 'FullColour', b + CARD_W - m - 0.17 * INCH, top + 0.01 * INCH, height=0.29 * INCH)
    # contact, bottom left
    c.setFillColor(cmyk('navy'))
    lines = [(PHONE, '  ·  24/7'), (EMAIL, ''), (ADDRESS, ''), (WEB, '')]
    y = b + m + 3 * 10.2
    for main, extra in lines:
        w = tracked(c, main, x, y, 'Montserrat-Medium', 7, 0, cmyk('navy'))
        if extra:
            tracked(c, extra, x + w, y, 'Montserrat-Medium', 7, 0, cmyk('coral_ink'))
        y -= 10.2


def build_cards():
    folder = '04_Business_Cards'
    # shared front
    path = p(folder, 'Print_Ready', 'Stagecoach_Card_Front.pdf')
    c, b = print_page(path, CARD_W, CARD_H)
    card_front(c, b)
    c.showPage()
    c.save()
    for name, role in TEAM:
        slug = name.split(',')[0].replace('.', '').replace(' ', '_')
        path = p(folder, 'Print_Ready', f'Stagecoach_Card_{slug}.pdf')
        c, b = print_page(path, CARD_W, CARD_H)
        card_front(c, b)
        c.showPage()
        card_back(c, b, name, role)
        c.showPage()
        c.save()
    # blank template for new hires
    path = p(folder, 'Print_Ready', 'Stagecoach_Card_TEMPLATE_Back.pdf')
    c, b = print_page(path, CARD_W, CARD_H)
    card_back(c, b, 'Full Name, Credentials', 'Job Title')
    c.showPage()
    c.save()


# ---------------------------------------------------------------- apparel
import base64  # noqa: E402


def font_css():
    out = []
    for w, n in ((400, 'Regular'), (500, 'Medium'), (600, 'SemiBold'), (700, 'Bold')):
        b64 = base64.b64encode(open(os.path.join(FONTS, f'Montserrat-{n}.ttf'), 'rb').read()).decode()
        out.append(f"@font-face{{font-family:'Montserrat';font-weight:{w};src:url(data:font/ttf;base64,{b64}) format('truetype')}}")
    return '\n'.join(out)


APPAREL = [
    # file stem, layout kind, tagline, axis, inches, colourways, note
    ('Tee_LeftChest_Print', 'horizontal', True, 'w', 3.5, ['Reverse', 'White', 'FullColour', 'Navy'], 'Screen print or DTG, left chest, 3.5 in wide'),
    ('Tee_LeftChest_Embroidery', 'horizontal', False, 'w', 3.5, ['Reverse', 'White', 'Navy'], 'Embroidery, left chest, 3.5 in wide. No tagline: it is too small to stitch'),
    ('Tee_Back_Print', 'stacked', True, 'w', 10.0, ['Reverse', 'White', 'FullColour', 'Navy'], 'Screen print or DTG, centred on the upper back, 10 in wide'),
    ('Cap_Front_Embroidery', 'symbol', True, 'h', 2.25, ['Reverse', 'White', 'Navy'], 'Embroidery, cap front, 2.25 in tall'),
]


def art_pdf(L, cw, size_pts, axis, path):
    """Artwork at exact print size, no clear space, transparent page."""
    k = size_pts / (L['w'] if axis == 'w' else L['h'])
    W, H = L['w'] * k, L['h'] * k
    c = rl_canvas.Canvas(path, pagesize=(W, H), initialFontName='Montserrat-Regular', initialFontSize=10)
    c.setTitle(os.path.splitext(os.path.basename(path))[0].replace('_', ' '))
    draw_logo(c, L, cw, 0, H, width=W)
    c.showPage()
    c.save()


TEE = ('M360 70 C400 118 600 118 640 70 L820 118 L985 330 L862 408 L782 322 L792 1062 '
       'C600 1078 400 1078 208 1062 L218 322 L138 408 L15 330 L180 118 Z')
COLLAR = 'M360 70 C400 118 600 118 640 70 C602 140 398 140 360 70 Z'


def tee_svg(fill, logo_svg, logo_x, logo_y, logo_w, back=False):
    shade = 'rgba(0,0,0,.18)' if fill != '#FFFFFF' and fill != '#F4F1E9' else 'rgba(0,0,0,.07)'
    collar = COLLAR if not back else 'M360 70 C400 92 600 92 640 70 C602 104 398 104 360 70 Z'
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="-20 0 1040 1100" width="520" height="550">'
            f'<defs><linearGradient id="g" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#fff" stop-opacity=".10"/>'
            f'<stop offset=".55" stop-color="#fff" stop-opacity="0"/><stop offset="1" stop-color="#000" stop-opacity=".12"/></linearGradient></defs>'
            f'<path d="{TEE}" fill="{fill}" stroke="{shade}" stroke-width="3"/>'
            f'<path d="{TEE}" fill="url(#g)"/>'
            f'<path d="{collar}" fill="{shade}"/>'
            f'<path d="M218 322 C236 380 236 420 230 470 M782 322 C764 380 764 420 770 470" stroke="{shade}" stroke-width="3" fill="none"/>'
            f'<g transform="translate({logo_x} {logo_y})">{logo_svg(logo_w)}</g></svg>')


def scaled_logo(L, cw):
    def f(w):
        k = w / L['w']
        return f'<g transform="scale({k:.5f})">{inner_svg(L, cw)}</g>'
    return f


def build_apparel():
    folder = '05_Apparel'
    for stem, kind, tag, axis, inches, cws, note in APPAREL:
        L = LG.layout(kind, tag)
        for cw in cws:
            base = f'Stagecoach_{stem}_{cw}'
            art_pdf(L, cw, inches * INCH, axis, p(folder, 'Artwork_Print_Ready', base + '.pdf'))
            svg = LG.to_svg(L, cw, pad=0)
            px = round(inches * 300)
            if axis == 'w':
                PNG.add(svg, p(folder, 'Artwork_PNG_300dpi', base + '.png'), width=px)
            else:
                PNG.add(svg, p(folder, 'Artwork_PNG_300dpi', base + '.png'), height=px)
    # mockups: a tee is about 20 in across the chest; the drawing's body is 570 units wide
    u = 570 / 20.0
    chest = LG.layout('horizontal', True)
    back = LG.layout('stacked', True)
    cap = LG.layout('symbol')
    scenes = []
    for fill, cw_front, cw_back, label in (('#1E2A3E', 'Reverse', 'Reverse', 'Navy tee'), ('#F4F1E9', 'FullColour', 'FullColour', 'Natural tee')):
        front = tee_svg(fill, scaled_logo(chest, cw_front), 560, 245, 3.5 * u)
        rear = tee_svg(fill, scaled_logo(back, cw_back), 500 - 10 * u / 2, 220, 10 * u, back=True)
        scenes.append((label, front, rear))
    css = font_css()
    for label, front, rear in scenes:
        slug = label.split()[0]
        html = (f'<html><head><style>{css} body{{margin:0;background:#E9E5DC;font-family:Montserrat;display:flex;gap:40px;'
                f'justify-content:center;align-items:center;height:100vh}} figure{{margin:0;text-align:center}}'
                f'figcaption{{font:600 15px Montserrat;letter-spacing:.18em;text-transform:uppercase;color:#1E2A3E;margin-top:14px}}</style></head>'
                f'<body><figure>{front}<figcaption>{label}, front</figcaption></figure><figure>{rear}<figcaption>{label}, back</figcaption></figure></body></html>')
        PNG.add_html(html, p(folder, 'Mockups', f'Mockup_{slug}_Tee.png'), 1200, 700, scale=2)
    # cap mockup
    k = 1.0
    cap_html = (f'<html><head><style>{css} body{{margin:0;background:#E9E5DC;display:grid;place-items:center;height:100vh}}</style></head><body>'
                f'<svg viewBox="0 0 600 420" width="600" height="420"><path d="M60 300 C60 120 180 60 300 60 C420 60 540 120 540 300 Z" fill="#1E2A3E"/>'
                f'<path d="M40 300 L560 300 C560 330 470 360 300 360 C130 360 40 330 40 300 Z" fill="#16202F"/>'
                f'<path d="M300 60 L300 300" stroke="rgba(0,0,0,.25)" stroke-width="3"/>'
                f'<circle cx="300" cy="64" r="9" fill="#16202F"/>'
                f'<g transform="translate({300 - cap["w"] * 0.95 / 2:.1f} 130)"><g transform="scale(.95)">{inner_svg(cap, "Reverse")}</g></g></svg>'
                f'</body></html>')
    PNG.add_html(cap_html, p(folder, 'Mockups', 'Mockup_Navy_Cap.png'), 700, 500, scale=2)
    open(p(folder, 'APPAREL_SPECS.txt'), 'w').write(apparel_specs())


def apparel_specs():
    lines = ['STAGECOACH APPAREL ARTWORK', '', 'Send the vendor the PDF from Artwork_Print_Ready. It is vector, CMYK, at exact print size.',
             'PNG files are 300 dpi at the same size, with a transparent background, for direct-to-garment printers that ask for PNG.', '']
    for stem, kind, tag, axis, inches, cws, note in APPAREL:
        lines.append(f'{stem.replace("_", " ")}: {note}.')
        lines.append(f'  Colourways: {", ".join(cws)}.')
    lines += ['',
              'Which colourway:',
              '  Navy or dark garments: Reverse (porcelain with a coral dot, 2 inks) or White (1 ink).',
              '  Natural, white or light garments: FullColour (navy with a coral dot, 2 inks) or Navy (1 ink).',
              '',
              'Ink colours: Navy #1E2A3E, Coral #E37970, Porcelain #F4F1E9. Ask the printer to match these on a strike-off before the run.',
              'Embroidery: the vendor digitises from the PDF. Thread: a navy close to #1E2A3E, an off-white close to #F4F1E9, a coral close to #E37970.',
              'Placement: left chest is about 4 in below the shoulder seam, centred on the wearer\'s left side. Back print sits about 3 in below the collar.']
    return '\n'.join(lines) + '\n'


# ---------------------------------------------------------------- stationery
LETTER_W, LETTER_H = 8.5 * INCH, 11 * INCH


def build_stationery():
    folder = '06_Stationery'
    # letterhead PDF, white stock, no bleed needed
    path = p(folder, 'Stagecoach_Letterhead.pdf')
    c = rl_canvas.Canvas(path, pagesize=(LETTER_W, LETTER_H), initialFontName='Montserrat-Regular', initialFontSize=10)
    c.setTitle('Stagecoach Letterhead')
    m = 0.85 * INCH
    draw_logo(c, LG.layout('horizontal', True), 'FullColour', m, LETTER_H - 0.75 * INCH, width=2.5 * INCH)
    y = 0.6 * INCH
    parts = [PHONE + ' · 24/7', EMAIL, ADDRESS, WEB]
    x = m
    for i, t in enumerate(parts):
        w = tracked(c, t, x, y, 'Montserrat-Medium', 7.5, 0, cmyk('navy'))
        x += w
        if i < len(parts) - 1:
            tracked(c, '   ·   ', x, y, 'Montserrat-Medium', 7.5, 0, cmyk('coral'))
            x += pdfmetrics.stringWidth('   ·   ', 'Montserrat-Medium', 7.5)
    c.showPage()
    c.save()
    # Word letterhead template
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    PNG.run()  # the header logo PNG must exist first
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin, sec.bottom_margin = Inches(1.45), Inches(1.0)
    sec.header_distance, sec.footer_distance = Inches(0.55), Inches(0.45)
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(os.path.join(OUT, '01_Logo', 'Horizontal', 'PNG', 'Stagecoach_Horizontal_FullColour_1600px.png'), width=Inches(2.85))
    fp = sec.footer.paragraphs[0]
    r = fp.add_run(f'{PHONE} · 24/7   ·   {EMAIL}   ·   {ADDRESS}   ·   {WEB}')
    r.font.name, r.font.size, r.font.color.rgb = 'Montserrat', Pt(7.5), RGBColor(0x1E, 0x2A, 0x3E)
    st = doc.styles['Normal']
    st.font.name, st.font.size, st.font.color.rgb = 'Montserrat', Pt(10.5), RGBColor(0x1E, 0x2A, 0x3E)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.25
    for t in ['Date', '', 'Recipient name', 'Organisation', 'Address', '', 'Dear Recipient,', '',
              'Write the letter here. This template uses Montserrat. If it is not installed, Word substitutes a similar font; install it from the Fonts folder in this kit.',
              '', 'Sincerely,', '', 'Your Name', 'Title, Stagecoach Concierge Emergency Medical Services']:
        doc.add_paragraph(t)
    doc.save(p(folder, 'Stagecoach_Letterhead_Template.docx'))
    # email signature
    sig = f'''<!-- Stagecoach email signature. Replace the name, title and the logo URL, then paste into your mail app's signature settings. -->
<table cellpadding="0" cellspacing="0" border="0" style="font-family:Montserrat,Arial,Helvetica,sans-serif;color:#1E2A3E;font-size:13px;line-height:1.5">
  <tr><td style="padding:0 0 10px 0"><img src="https://stagecoachmedical.com/brand/Stagecoach_Email_Logo.png" width="220" alt="Stagecoach Concierge Emergency Medical Services" style="display:block;width:220px;height:auto;border:0"></td></tr>
  <tr><td style="font-weight:700;font-size:15px">Full Name, Credentials</td></tr>
  <tr><td style="font-weight:600;font-size:11px;letter-spacing:1.5px;text-transform:uppercase;color:#A9473B;padding-bottom:8px">Job Title</td></tr>
  <tr><td><a href="tel:+12028399439" style="color:#1E2A3E;text-decoration:none">{PHONE}</a> <span style="color:#A9473B">· 24/7</span></td></tr>
  <tr><td><a href="mailto:{EMAIL}" style="color:#1E2A3E;text-decoration:none">{EMAIL}</a></td></tr>
  <tr><td><a href="https://{WEB}" style="color:#1E2A3E;text-decoration:none">{WEB}</a></td></tr>
  <tr><td style="font-size:11px;color:#5B6473;padding-top:10px">Not a substitute for 911. In a life-threatening emergency, call 911.</td></tr>
</table>
'''
    open(p(folder, 'Email_Signature', 'Stagecoach_Email_Signature.html'), 'w').write(sig)
    L = LG.layout('horizontal', True)
    PNG.add(LG.to_svg(L, 'FullColour', pad=0), p(folder, 'Email_Signature', 'Stagecoach_Email_Logo.png'), width=660)


# ---------------------------------------------------------------- digital
def build_digital():
    folder = '07_Digital'
    css = font_css()
    sym = LG.layout('symbol')
    hz = LG.layout('horizontal', True)

    def page(w, h, body, bg='#1E2A3E', glow=True):
        g = 'radial-gradient(55% 75% at 85% 15%,rgba(227,121,112,.32),transparent 65%),' if glow else ''
        return (f'<html><head><style>{css} body{{margin:0;width:{w}px;height:{h}px;overflow:hidden;position:relative;'
                f'background:{g}{bg};font-family:Montserrat;color:#F4F1E9}}</style></head><body>{body}</body></html>')

    def place(L, cw, x, y, w):
        k = w / L['w']
        return f'<svg style="position:absolute;left:{x}px;top:{y}px;overflow:visible" width="{w}" height="{L["h"] * k}"><g transform="scale({k:.5f})">{inner_svg(L, cw)}</g></svg>'

    # avatars, circle-safe
    for name, bg, cw in (('Navy', '#1E2A3E', 'Reverse'), ('Porcelain', '#F4F1E9', 'FullColour')):
        h = 520
        w = sym['w'] * h / sym['h']
        PNG.add_html(page(1080, 1080, place(sym, cw, (1080 - w) / 2, (1080 - h) / 2, w), bg=bg, glow=False),
                     p(folder, 'Social', f'Stagecoach_Avatar_{name}_1080.png'), 1080, 1080)
    line = 'Emergency Care That Comes to You.'
    sub = 'Washington, DC  ·  24/7  ·  202.839.9439'
    covers = [('LinkedIn_Cover_1584x396', 1584, 396, 560, 120), ('X_Header_1500x500', 1500, 500, 560, 150)]
    for name, w, h, lw, ytop in covers:
        lh = hz['h'] * lw / hz['w']
        body = (place(hz, 'Reverse', w - lw - 110, (h - lh) / 2 - 26, lw)
                + f'<div style="position:absolute;right:110px;top:{(h + lh) / 2 + 2}px;font:500 20px Montserrat;letter-spacing:.06em;opacity:.85">{sub}</div>')
        PNG.add_html(page(w, h, body), p(folder, 'Social', f'Stagecoach_{name}.png'), w, h)
    # video call background, logo bottom left, calm space for the person
    lw = 420
    lh = hz['h'] * lw / hz['w']
    PNG.add_html(page(1920, 1080, place(hz, 'Reverse', 96, 1080 - 96 - lh, lw)), p(folder, 'Video_Call', 'Stagecoach_Video_Background_1920x1080.png'), 1920, 1080)
    # link preview card from the site
    shutil.copy(os.path.join(SITE, 'assets', 'og.jpg'), p(folder, 'Web', 'Stagecoach_Link_Preview_1200x630.jpg'))
    for f in ('favicon.svg', 'favicon-32.png', 'apple-touch-icon.png'):
        shutil.copy(os.path.join(SITE, 'assets', f), p(folder, 'Web', f))


# ---------------------------------------------------------------- colour + type
def build_colour_type():
    folder = '02_Colour'
    order = ['navy', 'night', 'coral', 'peach', 'cream', 'coral_ink']
    rows = ['STAGECOACH COLOUR', '', 'HEX is the master value. CMYK is a starting build for coated paper: approve a printed proof against the HEX swatch before a full run.',
            'Pantone: ask the printer to match the HEX on their Pantone guide and note the chosen numbers here for reuse.', '']
    for k in order:
        c = LG.COLOURS[k]
        rows.append(f"{c['name']:<16} HEX {c['hex']}   RGB {c['rgb'][0]} {c['rgb'][1]} {c['rgb'][2]}   CMYK {' '.join(str(v) for v in c['cmyk'])}")
    rows += ['', 'Use:', '  Navy and Porcelain carry the brand. Coral is an accent: the dot in the S, small highlights, the waitlist button.',
             '  Deep Coral replaces Coral for small text on light backgrounds, so it stays readable.', '  Night is for deep backgrounds and gradients behind Navy.']
    open(p(folder, 'Stagecoach_Colours.txt'), 'w').write('\n'.join(rows) + '\n')
    css = ':root{\n' + ''.join(f"  --sc-{k.replace('_', '-')}: {LG.COLOURS[k]['hex']};\n" for k in order) + '}\n'
    open(p(folder, 'Stagecoach_Colours.css'), 'w').write(css)
    write_ase(p(folder, 'Stagecoach_Colours_RGB.ase'), order, 'rgb')
    write_ase(p(folder, 'Stagecoach_Colours_CMYK.ase'), order, 'cmyk')
    tf = '03_Typography'
    for f in ('Montserrat-Regular.ttf', 'Montserrat-Medium.ttf', 'Montserrat-SemiBold.ttf', 'Montserrat-Bold.ttf', 'OFL.txt'):
        shutil.copy(os.path.join(FONTS, f), p(tf, 'Montserrat', f))
    open(p(tf, 'TYPE_GUIDE.txt'), 'w').write(
        'STAGECOACH TYPE\n\nOne family: Montserrat. Free under the SIL Open Font License (OFL.txt), so it can be installed on any computer and given to vendors.\n'
        'Install: open each .ttf file and click Install.\n\n'
        'Headlines: Montserrat Medium, tight letter spacing (-2% to -3%), Title Case.\n'
        'Labels and small capitals: Montserrat SemiBold, all caps, letter spacing +14% to +20%.\n'
        'Body text: Montserrat Regular or Medium, 1.5 line spacing.\n'
        'Names on cards and signatures: Montserrat Bold.\n\n'
        'Do not use a serif font or any second family. The logo itself is outlined artwork: never retype it.\n')


def write_ase(path, keys, model):
    import struct

    def block(name, values):
        n = (name + '\0').encode('utf-16-be')
        body = struct.pack('>H', len(name) + 1) + n
        if model == 'rgb':
            body += b'RGB ' + struct.pack('>fff', *values)
        else:
            body += b'CMYK' + struct.pack('>ffff', *values)
        body += struct.pack('>H', 2)
        return struct.pack('>HI', 0x0001, len(body)) + body

    blocks = []
    for k in keys:
        c = LG.COLOURS[k]
        vals = [v / 255 for v in c['rgb']] if model == 'rgb' else [v / 100 for v in c['cmyk']]
        blocks.append(block(c['name'] if c['name'].startswith('Stagecoach') else f"Stagecoach {c['name']}", vals))
    open(path, 'wb').write(b'ASEF' + struct.pack('>HHI', 1, 0, len(blocks)) + b''.join(blocks))


START_HERE = '''STAGECOACH BRAND KIT
Stagecoach Concierge Emergency Medical Services

Start with Stagecoach_Brand_Vision.pdf for the full picture. This note says which file to send to whom.

WHICH FILE DO I NEED?

Printer (cards, flyers, signage, anything on paper)
  Logos: 01_Logo > [layout] > PDF_print or EPS_print. Vector, CMYK, outlined, no fonts needed.
  Business cards: 04_Business_Cards > Print_Ready. One 2-page PDF per person (front, back),
  3.5 x 2 in with 0.125 in bleed, trim and bleed boxes set. Ask for a printed proof first.

Apparel and merchandise vendor
  05_Apparel > Artwork_Print_Ready (vector PDF at exact print size) and APPAREL_SPECS.txt.
  Embroiderers digitise from the same PDF.

Website, social media, slides, documents
  01_Logo > [layout] > PNG (transparent) or SVG. 07_Digital for profile pictures and covers.

Microsoft Word or Google Docs
  06_Stationery > Stagecoach_Letterhead_Template.docx. Install the fonts in 03_Typography first.

Email signature
  06_Stationery > Email_Signature. Upload Stagecoach_Email_Logo.png somewhere public
  (for example the website), put its address in the HTML, then paste into your mail app.

Designers
  02_Colour has Adobe swatch files (.ase) and CSS. 03_Typography has Montserrat.

WHICH LOGO?

  Horizontal            Everyday use. Website header, letterhead, documents.
  Horizontal_NoTagline  Anything under about 2 in or 200 px wide, and all embroidery.
  Stacked               Square or tall spaces: tote bags, back prints, social posts.
  Symbol                Very small spaces: app icon, favicon, caps, embossing.
  Wordmark              Rare. Only where the symbol already appears nearby.

WHICH COLOUR VERSION?

  FullColour  Navy with a coral dot, on light backgrounds.
  Reverse     Porcelain with a coral dot, on navy or dark backgrounds.
  Navy, White, Black  One colour, for single-ink printing, foil, embossing and engraving.

RULES

  Keep clear space of half the S height on every side.
  Never stretch, rotate, recolour, add effects, or retype the logo.
  Do not place the logo on busy photographs.
  One font family only: Montserrat.

BEFORE PRINTING ANYTHING

  CMYK values are a starting point. Approve a printed proof against the HEX colours.
  Business cards use the shared phone and email. Confirm each person's details,
  and whether anyone should have a personal email or mobile, before ordering.
  Card names drop "Dr." where MD or DDS follows (for example David Vitberg, MD). Confirm.
'''


def build_readme():
    open(p('00_START_HERE.txt'), 'w').write(START_HERE)


if __name__ == '__main__':
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)
    build_logos()
    build_cards()
    build_colour_type()
    build_apparel()
    build_stationery()
    build_digital()
    build_readme()
    PNG.run()
    n = sum(len(fs) for _, _, fs in os.walk(OUT))
    print(f'wrote {OUT} ({n} files)')
